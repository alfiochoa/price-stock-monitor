from pathlib import Path

import pytest
import requests

from pricemon import mockstore
from pricemon.alerts import Alerter
from pricemon.cli import run_once
from pricemon.config import load_site, load_sites
from pricemon.fetch import Fetcher, FetchError, RobotsDisallowed
from pricemon.normalize import parse_availability, parse_price
from pricemon.scrape import Item, parse_page, scrape_site
from pricemon.store import Store

FIX = Path(__file__).parent / "fixtures"
SITES = Path(__file__).parent.parent / "sites"


# ---------- normalisation ----------
@pytest.mark.parametrize("raw,exp", [
    ("£51.77", (51.77, "GBP")), ("$1,299.00", (1299.0, "USD")), ("1.299,00 €", (1299.0, "EUR")),
    ("€14,50", (14.5, "EUR")), ("USD 20", (20.0, "USD")), ("free", (None, None)), (None, (None, None)),
])
def test_parse_price(raw, exp):
    assert parse_price(raw) == exp


def test_parse_price_default_currency():
    assert parse_price("19.99", "EUR") == (19.99, "EUR")


@pytest.mark.parametrize("raw,exp", [
    ("In stock", "in_stock"), ("In stock (22 available)", "in_stock"), ("Sold out", "out_of_stock"),
    ("Currently unavailable", "out_of_stock"), ("Not available", "out_of_stock"), ("", "unknown"), ("ships soon", "unknown"),
])
def test_parse_availability(raw, exp):
    assert parse_availability(raw) == exp


# ---------- parsing saved HTML fixtures ----------
def test_books_fixture():
    cfg = load_site(SITES / "books_toscrape.yaml")
    items, filled, containers, nxt = parse_page((FIX / "books_page2.html").read_text(), cfg,
                                                "https://books.toscrape.com/catalogue/page-2.html")
    assert containers == len(items) == 20
    assert items[0].name == "In Her Wake" and items[0].sku == "in-her-wake_980"
    assert items[0].price == 12.84 and items[0].currency == "GBP" and items[0].availability == "in_stock"
    assert nxt == "https://books.toscrape.com/catalogue/page-3.html"


def test_webscraper_static_fixture():
    cfg = load_site(SITES / "webscraper_static.yaml")
    url = "https://webscraper.io/test-sites/e-commerce/static/computers/laptops"
    items, _, _, nxt = parse_page((FIX / "webscraper_laptops.html").read_text(), cfg, url)
    assert items and items[0].sku == "31" and items[0].price == 416.99 and items[0].currency == "USD"
    assert items[0].availability == "unknown"
    assert nxt.endswith("laptops?page=2")


def test_webscraper_js_rendered_fixture():
    cfg = load_site(SITES / "webscraper_js.yaml")
    items, filled, containers, _ = parse_page((FIX / "webscraper_ajax_rendered.html").read_text(), cfg, "https://webscraper.io/x")
    assert containers > 0 and filled["price"] == containers and items[0].price is not None


# ---------- breakage detector ----------
class FakeResp:
    def __init__(self, status=200, text="", headers=None):
        self.status_code, self.text, self.headers = status, text, headers or {}

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(f"{self.status_code}")


class FakeSession:
    """Serves {url: [responses...]}; robots.txt is permissive unless provided."""
    def __init__(self, pages):
        self.pages, self.headers, self.calls = pages, {}, []

    def get(self, url, timeout=None):
        self.calls.append(url)
        if url.endswith("/robots.txt"):
            return self.pages.get(url, FakeResp(404))
        seq = self.pages[url]
        r = seq.pop(0) if len(seq) > 1 else seq[0]
        if isinstance(r, Exception):
            raise r
        return r


def fetcher(pages, **kw):
    return Fetcher(cache_dir=None, session=FakeSession(pages), sleep=lambda s: None, **kw)


def test_breakage_when_item_selector_matches_nothing():
    cfg = load_site(SITES / "books_toscrape.yaml")
    url = "https://books.toscrape.com/catalogue/page-1.html"
    res = scrape_site(cfg, fetcher({url: [FakeResp(text="<html><body>redesigned</body></html>")]}))
    assert res.items == [] and "matched 0 elements" in res.breakages[0]


def test_breakage_when_price_selector_silently_empty():
    cfg = load_site(SITES / "books_toscrape.yaml")
    html = (FIX / "books_page2.html").read_text().replace("price_color", "price_new")
    res = scrape_site(cfg, fetcher({"https://books.toscrape.com/catalogue/page-1.html": [FakeResp(text=html)]}))
    assert any("field 'price'" in b for b in res.breakages)


def test_healthy_page_has_no_breakage():
    cfg = load_site(SITES / "books_toscrape.yaml")
    html = (FIX / "books_page2.html").read_text()
    cfg.max_pages = 1
    res = scrape_site(cfg, fetcher({"https://books.toscrape.com/catalogue/page-1.html": [FakeResp(text=html)]}))
    assert res.breakages == [] and len(res.items) == 20


# ---------- fetcher: retries, robots, cache, rate limit ----------
def test_retries_with_exponential_backoff():
    sleeps = []
    f = Fetcher(cache_dir=None, sleep=sleeps.append, backoff=1.0, rate_limit=0,
                session=FakeSession({"http://x.test/a": [FakeResp(503), requests.ConnectionError("boom"), FakeResp(200, "ok")]}))
    assert f.get("http://x.test/a") == "ok"
    assert sleeps == [1.0, 2.0] and f.stats["retries"] == 2


def test_gives_up_after_retries():
    f = fetcher({"http://x.test/a": [FakeResp(500)]}, retries=2)
    with pytest.raises(FetchError):
        f.get("http://x.test/a")


def test_404_is_not_retried():
    f = fetcher({"http://x.test/a": [FakeResp(404)]})
    with pytest.raises(FetchError):
        f.get("http://x.test/a")
    assert f.stats["retries"] == 0


def test_robots_disallow_respected():
    robots = FakeResp(200, "User-agent: *\nDisallow: /private/\n")
    f = fetcher({"http://x.test/robots.txt": robots, "http://x.test/private/a": [FakeResp(200, "secret")]})
    with pytest.raises(RobotsDisallowed):
        f.get("http://x.test/private/a")
    assert "http://x.test/private/a" not in f.session.calls


def test_cache_avoids_second_request(tmp_path):
    f = Fetcher(cache_dir=tmp_path, sleep=lambda s: None, session=FakeSession({"http://x.test/a": [FakeResp(200, "hi")]}))
    assert f.get("http://x.test/a") == f.get("http://x.test/a") == "hi"
    assert f.stats == {"network": 1, "cache_hits": 1, "retries": 0}


def test_rate_limit_sleeps_between_requests():
    sleeps, t = [], [0.0]
    f = Fetcher(cache_dir=None, rate_limit=2.0, sleep=sleeps.append, clock=lambda: t[0],
                session=FakeSession({"http://x.test/a": [FakeResp(200, "a")], "http://x.test/b": [FakeResp(200, "b")]}))
    f.get("http://x.test/a")
    t[0] = 0.5
    f.get("http://x.test/b")
    assert sleeps == [1.5]


# ---------- store: change detection ----------
def item(price, avail="in_stock", sku="A"):
    return Item(sku, "Widget", price, "EUR", avail)


def test_change_detection_kinds():
    s = Store(":memory:")
    assert s.record("t", [item(100.0)], ts="2026-01-01") == []
    ev = s.record("t", [item(90.0)], ts="2026-01-02")
    assert [(e.kind, e.pct) for e in ev] == [("price_drop", -10.0)]
    ev = s.record("t", [item(99.0, "out_of_stock")], ts="2026-01-03")
    assert [e.kind for e in ev] == ["price_rise", "out_of_stock"]
    ev = s.record("t", [item(99.0, "in_stock")], ts="2026-01-04")
    assert [e.kind for e in ev] == ["back_in_stock"]
    assert s.record("t", [item(99.0)], ts="2026-01-05") == []  # unchanged -> silent
    assert s.db.execute("SELECT COUNT(*) FROM observations").fetchone()[0] == 5


def test_min_pct_threshold_and_unknown_availability():
    s = Store(":memory:")
    s.record("t", [item(100.0)], ts="2026-01-01")
    assert s.record("t", [item(99.5, "unknown")], ts="2026-01-02") == []  # 0.5% < 1%, unknown never flips stock


# ---------- end to end against the mock store ----------
def test_end_to_end_mock_store_with_breakage_day(tmp_path):
    srv = mockstore.MockServer(port=0).start()
    try:
        cfg = load_sites(SITES)["mock_store"]
        cfg.base_url = f"http://127.0.0.1:{srv.port}"
        store, alerter = Store(":memory:"), Alerter(url="", log_path=None)
        f = lambda: Fetcher(cache_dir=None, rate_limit=0, retries=0)
        srv.set_day(0)
        run_once([cfg], store, f(), alerter, ts="2026-01-01")
        srv.set_day(1)
        s = run_once([cfg], store, f(), alerter, ts="2026-01-02")
        assert s["mock_store"]["items"] == 12 and s["mock_store"]["events"] == 2
        assert any("Trail Runner" in m and "-10.0%" in m for m in alerter.sent)
        srv.set_day(mockstore.BREAKAGE_DAY)
        s = run_once([cfg], store, f(), alerter, ts="2026-01-03")
        assert s["mock_store"]["breakages"] == 1 and any("breakage" in m for m in alerter.sent)
    finally:
        srv.stop()


def test_alerter_posts_slack_and_discord_payloads(monkeypatch):
    posted = []
    monkeypatch.setattr("pricemon.alerts.requests.post",
                        lambda url, json, timeout: posted.append((url, json)) or type("R", (), {"raise_for_status": lambda s: None})())
    Alerter(url="http://hook", kind="slack").send("hi")
    Alerter(url="http://hook", kind="discord").send("yo")
    assert posted == [("http://hook", {"text": "hi"}), ("http://hook", {"content": "yo"})]


def test_export_csv_and_xlsx(tmp_path):
    from pricemon.export import export
    s = Store(":memory:")
    s.record("t", [item(10.0)], ts="2026-01-01")
    s.record("t", [item(12.0)], ts="2026-01-02")
    assert "12.0" in Path(export(s, str(tmp_path / "new" / "o.csv"))).read_text()
    import openpyxl
    wb = openpyxl.load_workbook(export(s, str(tmp_path / "o.xlsx")))
    assert wb.sheetnames == ["latest", "history", "events"] and wb["history"].max_row == 3


def test_seed_demo_and_streamlit_app(tmp_path):
    from streamlit.testing.v1 import AppTest

    from pricemon.seed import seed_demo
    db = seed_demo(str(tmp_path / "d.db"))
    s = Store(db)
    assert s.db.execute("SELECT COUNT(*) FROM observations").fetchone()[0] == 72
    assert s.db.execute("SELECT COUNT(*) FROM events").fetchone()[0] == 10
    assert s.db.execute("SELECT SUM(breakages) FROM runs").fetchone()[0] == 2
    import os
    os.environ["PRICEMON_DB"] = db
    try:
        at = AppTest.from_file("../streamlit_app.py", default_timeout=60).run()
    finally:
        del os.environ["PRICEMON_DB"]
    assert not at.exception and at.metric[0].value == "12"
